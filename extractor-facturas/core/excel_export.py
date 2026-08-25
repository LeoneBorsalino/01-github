"""Construcción del archivo Excel de salida a partir de las facturas procesadas."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .parser import FacturaExtraida

ENCABEZADO_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
ENCABEZADO_FONT = Font(color="FFFFFF", bold=True)
ESTADO_OK_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
ESTADO_REVISAR_FILL = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
ESTADO_ERROR_FILL = PatternFill(start_color="FECACA", end_color="FECACA", fill_type="solid")

COLUMNAS_FACTURAS = [
    ("Archivo", 28),
    ("Estado", 12),
    ("Proveedor", 28),
    ("CUIT/RFC Proveedor", 18),
    ("Cliente", 24),
    ("CUIT/RFC Cliente", 18),
    ("Condición IVA", 20),
    ("N° Factura", 16),
    ("Fecha Emisión", 14),
    ("Fecha Vencimiento", 16),
    ("Moneda", 10),
    ("Subtotal", 14),
    ("IVA", 12),
    ("Total", 14),
    ("Método extracción", 16),
    ("Observaciones", 30),
]

COLUMNAS_ITEMS = [
    ("Archivo", 28),
    ("N° Factura", 16),
    ("Descripción", 40),
    ("Cantidad", 10),
    ("Precio Unitario", 14),
    ("Importe", 14),
]


def _escribir_encabezado(ws: Worksheet, columnas: list[tuple[str, int]]) -> None:
    for col_idx, (titulo, ancho) in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=col_idx, value=titulo)
        celda.fill = ENCABEZADO_FILL
        celda.font = ENCABEZADO_FONT
        celda.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho
    ws.freeze_panes = "A2"


def _fill_estado(estado: str) -> PatternFill:
    return {
        "OK": ESTADO_OK_FILL,
        "Revisar": ESTADO_REVISAR_FILL,
        "ERROR": ESTADO_ERROR_FILL,
    }.get(estado, ESTADO_REVISAR_FILL)


def generar_excel(facturas: list[FacturaExtraida]) -> bytes:
    wb = Workbook()

    ws_facturas = wb.active
    ws_facturas.title = "Facturas"
    _escribir_encabezado(ws_facturas, COLUMNAS_FACTURAS)

    ws_items = wb.create_sheet("Items")
    _escribir_encabezado(ws_items, COLUMNAS_ITEMS)

    ws_texto = wb.create_sheet("Texto OCR (verificación)")
    _escribir_encabezado(ws_texto, [("Archivo", 28), ("Método", 14), ("Texto extraído", 100)])

    fila_f = 2
    fila_i = 2
    fila_t = 2

    for f in facturas:
        observaciones = f.error if f.error else ", ".join(f.campos_faltantes())
        fila = [
            f.archivo,
            f.estado,
            f.proveedor,
            f.cuit_proveedor,
            f.cliente,
            f.cuit_cliente,
            f.condicion_iva,
            f.numero_factura,
            f.fecha_emision,
            f.fecha_vencimiento,
            f.moneda,
            f.subtotal,
            f.iva,
            f.total,
            "Texto PDF" if f.metodo == "texto_pdf" else "OCR",
            observaciones,
        ]
        for col_idx, valor in enumerate(fila, start=1):
            celda = ws_facturas.cell(row=fila_f, column=col_idx, value=valor)
            if col_idx in (9, 10) and valor is not None:
                celda.number_format = "DD/MM/YYYY"
            if col_idx in (12, 13, 14) and valor is not None:
                celda.number_format = "#,##0.00"
        ws_facturas.cell(row=fila_f, column=2).fill = _fill_estado(f.estado)
        fila_f += 1

        for item in f.items:
            fila_item = [
                f.archivo,
                f.numero_factura,
                item.get("descripcion"),
                item.get("cantidad"),
                item.get("precio_unitario"),
                item.get("importe"),
            ]
            for col_idx, valor in enumerate(fila_item, start=1):
                celda = ws_items.cell(row=fila_i, column=col_idx, value=valor)
                if col_idx in (5, 6) and valor is not None:
                    celda.number_format = "#,##0.00"
            fila_i += 1

        celda_texto = ws_texto.cell(row=fila_t, column=1, value=f.archivo)
        ws_texto.cell(row=fila_t, column=2, value="Texto PDF" if f.metodo == "texto_pdf" else "OCR")
        texto_recortado = (f.texto_crudo or "")[:32000]
        ws_texto.cell(row=fila_t, column=3, value=texto_recortado).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        fila_t += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
